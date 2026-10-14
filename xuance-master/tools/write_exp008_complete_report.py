"""Build the Chinese exp008 technical report from frozen records.

Run with the Codex bundled Python. Plotting is a separate preparatory command.
This script does not import xuance, modify training configuration, or start jobs.
"""
import csv
import hashlib
import json
from collections import Counter
from datetime import datetime, timezone, timedelta
from pathlib import Path
from xml.sax.saxutils import escape
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / 'docs'
RESULT = ROOT / 'outputs/results/matd3/baseline_exp008'
REVIEW = ROOT / 'outputs/reviews/exp008_final_v1'
STEM = 'baseline_exp008_complete_report_2026-10-13'
blocks = []

def read(path):
    return json.loads((ROOT / path).read_text(encoding='utf-8-sig'))

def csvrows(name):
    with (RESULT / name).open(encoding='utf-8-sig', newline='') as stream:
        return list(csv.DictReader(stream))

def h(text, level=1): blocks.append(('h', level, text))
def p(text): blocks.append(('p', text))
def table(headers, rows, widths=None): blocks.append(('table', headers, rows, widths))
def picture(name, caption, width=6.7): blocks.append(('image', name, caption, width))
def code(text): blocks.append(('code', text))
def link(label, path): blocks.append(('link', label, str((ROOT / path).resolve()).replace('\\', '/')))

# Real structural MathML is transformed to editable Office Math, rather than
# writing raw TeX or flattening fractions/subscripts into a plain-text run.
def mi(t): return '<mi>'+escape(str(t))+'</mi>'
def mn(t): return '<mn>'+escape(str(t))+'</mn>'
def mo(t): return '<mo>'+escape(str(t))+'</mo>'
def mt(t): return '<mtext>'+escape(str(t))+'</mtext>'
def row(*parts): return '<mrow>'+''.join(parts)+'</mrow>'
def sub(a,b): return '<msub>'+a+b+'</msub>'
def sup(a,b): return '<msup>'+a+b+'</msup>'
def frac(a,b): return '<mfrac>'+a+b+'</mfrac>'
def par(a): return row(mo('('),a,mo(')'))
def eq(tex, mml): blocks.append(('eq', tex, mml))
def S(a,b): return sub(mi(a),mi(b))

config = read('outputs/results/matd3/baseline_exp008/config.json')
meta = read('outputs/experiments/exp008/metadata.json')
process = read('outputs/experiments/exp008/process.json')
summary = read('outputs/results/matd3/baseline_exp008/baseline_summary.json')
manifest = read('outputs/reviews/exp008_final_v1/manifest.json')
review_process = read('outputs/experiments/exp008/final_review_process.json')
retention_process = read('outputs/reviews/exp008_final_v1/retention_process.json')
episodes = csvrows('training_episodes.csv')
evaluations = csvrows('curriculum_eval.csv')
updates = csvrows('training_updates.csv')
current = [r for r in evaluations if r['phase'] == 'current']
finals = [read(f'outputs/reviews/exp008_final_v1/final008_level{i}.json') for i in range(1,6)]
initials = [read(f'outputs/reviews/exp008_final_v1/initial007_level{i}.json') for i in range(1,6)]
retained = [read(f'outputs/reviews/exp008_final_v1/beststage_level{i}.json') for i in range(1,4)]
videos = []
for i in range(1,6):
    folder = 'exp008_stage1_video_v1' if i == 1 else 'exp008_stages2to5_video_v1'
    relative = f'outputs/reviews/{folder}/level_{i}'
    videos.append((relative, read(relative+'/manifest.json')))

def pct(v): return f'{float(v)*100:.0f}%'
def china(t): return datetime.fromisoformat(t).astimezone(timezone(timedelta(hours=8))).strftime('%Y-%m-%d %H:%M:%S')

assert process['status'] == 'completed' and process['exit_code'] == 0
assert review_process['status'] == 'completed' and review_process['exit_code'] == 0
assert retention_process['status'] == 'completed' and retention_process['exit_code'] == 0
assert summary['environment_steps'] == 3000000
assert manifest['source_hash_count'] == 29 and not manifest['source_hash_mismatches']
assert int(updates[-1]['updates']) == 187244 and int(updates[-1]['actor_updates']) == 93622
for report in finals+initials+retained:
    assert report['metrics']['episodes'] == 100 and report['env_seed'] == 60001
    assert report['exploration_noise'] is False and report['collision_termination'] is False

blocks.append(('title', '无人机追捕实验008设置与结果总结'))
p('完整技术记录与后续工作说明    2026年10月13日    分支 feat-coverage-reaward')
p('本报告用于理解和复现实验008：说明任务、固定城市、初始状态、三维运动学、逃逸机势场、网络观测、MATD3训练、课程升级、实际奖励及日志含义，并给出最终结果与下一轮待做事项。事实以本轮冻结配置、源文件、完整CSV和模型复评为准。')
p('实验008已正常完成300万环境转移，并进入第5级；第5级未达到安全成功率至少80%且碰撞率至多10%的通过条件。训练末次第5级评估为捕获87%、安全成功65%、碰撞35%；冻结最终模型在新种子60001下复评为捕获80%、安全成功57%、碰撞43%。')
p('更需要处理的问题是旧级能力保留。相同60001种子下，第1至3级历史最好模型的安全成功率分别为100%、99%、100%，最终模型则为38%、4%、6%。这一对照支持后续课程训练中发生旧级安全能力遗忘的判断。当前环境能够训练出部分有效的闭环追捕与避障行为，但证据尚不足以说明最终单一模型掌握了五级稳定安全追捕，也不能据此宣称具备真实无人机部署能力。')
p('你已采纳的物理步0.1秒加决策步0.5秒列入下一轮待做。008实际仍使用物理步0.1秒、决策步1秒。本报告没有把未实施的调整写成既有结果。自动监控按你的要求保持暂停，后续按需查看。')

h('1 实验范围与结果口径')
h('1 1 运行来源和结束状态',2)
table(['项目','已核实值'],[
    ['实验与算法','008  MATD3  四架追捕机共享参数'],
    ['Git分支','feat-coverage-reaward'],
    ['Git基准提交',meta['git_head']],
    ['开始时间',china(process['started_at'])+' 北京时间'],
    ['结束时间',china(process['finished_at'])+' 北京时间'],
    ['训练预算和实际转移','3,000,000  已全部完成'],
    ['训练并行环境','16  每次向量交互通常增加16个环境转移'],
    ['优化更新','Critic 187,244次  Actor 93,622次'],
    ['训练摘要耗时',f"{summary['elapsed_seconds']:,.3f}秒  约3小时6分14秒"],
    ['退出状态','监督记录completed  真实退出码0  原训练进程已不存在'],
    ['训练源核验','metadata记录的29项源文件哈希在复评前全部一致'],
    ['初始化方式','007最终模型完整策略权重初始化  不是随机权重从头训练'],
], [1.7,5.0])
p('服务器独立工作树为 /home/ryy/UAV_Mine_baseline_coverage/xuance-master；本机镜像为 baseline_work。原 /home/ryy/UAV_Mine 工作树及用户已有修改保留。当前工作未提交、未推送。新增的总结、复评和渲染脚本是交付工具，不属于本轮开始前冻结的29项训练源。')
p('最终总步数来自正常结束的baseline_summary.json，不能用略早的training_status.json代替。末次优化CSV记录在2,999,984转移，实际计数为187,244和93,622；该计数涵盖末次写出窗口。一次Critic更新同时更新两张Q网络，不能把计数再乘二。')
h('1 2 三种成绩必须分别解释',2)
p('历史最好成绩描述某一级训练时保存的best_stage模型；训练末次成绩描述当时正在训练课程的最后一次评估；最终模型逐级复评描述同一最终权重在五级上的实际能力。训练阶段曾达到100%安全成功，不代表后续训练结束仍保留该能力。')
p('训练期间正式评估使用10001种子，按同一评估设置反复选择课程和模型。结束后用未用于008训练与选模的60001种子对最终模型做五级复评，每级100回合、8个并行环境、无探索、最多400决策步、统计完整实际回合碰撞。60001仍使用同一套固定城市，只更换出生和随机航向序列；它不构成未见城市布局测试。')

h('2 追捕任务与终止判据')
p('每个环境包含四架学习策略控制的追捕机和一架脚本控制的逃逸机。目标是在三维城市空域中安全接近目标并满足捕获判据。所有追捕机共享一次团队捕获奖励，个体接近、建筑和边界奖励分别计算；机间碰撞按你的要求不考虑。')
table(['条件','具体判定与作用'],[
    ['单机直接拦截','任意一架追捕机与目标的三维距离进入5米范围即捕获'],
    ['单机连续保持','同一架追捕机在当前课程捕获半径内连续保持2秒即捕获'],
    ['保持计时','各机单独累计  出界则相应计时重置  不能四架轮流拼接2秒'],
    ['训练碰撞终止','任意追捕机出现建筑或边界真实接触  全队回合提前结束'],
    ['超时截断','达到400决策步  008最多400秒  不额外发放全队失败奖励'],
    ['机间接触','不计碰撞  不提前终止  不产生机间避障奖励'],
    ['围捕几何','没有要求多机包围角度  没有覆盖比例门槛  没有视线可见性条件'],
], [1.6,5.1])
p('距离使用三维欧氏距离。捕获检查利用同一物理子步内的相对运动线段，计算进入半径的区间和连续保持时间；捕获发生在子步中间时，记录分数时刻并插值回退位置、速度与加速度。因而捕获时间可以是64.30秒，而不是整数秒。')
p('训练中若同一决策步同时记录到碰撞和捕获，原始捕获标记可能为真，但安全成功为假；失败分支不发捕获奖励。统计时应同时检查safe_success及collision，而不能只看success。')
h('2 1 评估中的安全成功',2)
p('正式评估关闭建筑和边界的碰撞提前终止，让撞过的策略继续执行到捕获或超时，并累计该实际回合内的接触。因此安全成功率是“完成捕获且全程没有追捕机接触”的回合比例。发生一次接触后即使后来捕获，也只能计入捕获率，不能计入安全成功率。捕获已经结束的回合不会为了补满400步继续运行。')
p('建筑碰撞率与边界碰撞率可以在同一回合重叠，总碰撞率表示至少一种接触，二者不能直接相加。训练关闭机间冲突后，日志中的机间碰撞率和r_teammate为0；最小机间距离仍可作为诊断量。逃逸机墙面接触单独记录，不等同于追捕机安全成功判据。')

h('3 固定城市环境')
h('3 1 地图尺度和建筑类型',2)
p('地图水平范围为1000米乘1000米，允许高度为绝对z坐标50至350米。机体简化为半径0.5米的几何体。建筑数据为轴对齐柱体，以xmin、xmax、ymin、ymax、建筑顶部高度五个数保存；建筑从地面向上延伸。允许水平绕行，也允许飞越楼顶。')
table(['类型','比例','平面每条边的候选尺寸','顶部高度'],[
    ['普通建筑','50%','20米或40米','10至60米'],
    ['高层建筑','40%','40米或60米','60至150米'],
    ['地标建筑','10%','60米或80米','150至350米'],
], [1.2,.7,3.0,1.8])
p('宽和长分别从对应候选边长中按固定组合取值，并不要求所有建筑都是正方形。高度由生成脚本在各类别范围内确定性分布，并按5米刻度取整。这五张地图是直接构建并保存的固定环境，不采用每回合重新随机布局，也不采用首次随机生成后固定。')
h('3 2 五张地图的均匀与递增构造',2)
p('tools/generate_uniform_city_maps.py在10乘10的网格上建立100个候选中心，横纵坐标均为50加100乘网格索引。按(row加2乘col)模5的分组顺序逐批加入20栋，使每一级在各行各列均匀分布。20栋地图是40栋的子集，继续嵌套到100栋，已存在的建筑位置、尺寸和高度保持一致。每加入20栋，严格含10栋普通、8栋高层、2栋地标。')
table(['课程','数量','普通 高层 地标','平面占地比例','顶部低于50米'],[
    ['1','20','10  8  2','3.68%','7'],
    ['2','40','20  16  4','7.52%','14'],
    ['3','60','30  24  6','11.60%','21'],
    ['4','80','40  32  8','15.52%','28'],
    ['5','100','50  40  10','19.20%','35'],
], [.55,.6,2.1,1.55,1.9])
p('最大平面边长80米，候选中心相隔100米，相邻建筑最小平面间隙约20米。占地比例描述二维覆盖，不代表整个可飞空域的三维占用率。顶部低于50米的普通楼体基本位于最低飞行高度下方；在180至220米出生时，普通建筑和绝大多数高层建筑可直接飞越，地标建筑更可能影响轨迹。因此不能仅凭100栋数量判断这是密集低空街区穿行任务。')
p('固定且嵌套的地图降低了学习环境变化，适合建立最简单基线；同时也使成绩容易依赖已见布局。当前环境还没有验证对新街区、旋转城市或新建筑排列的泛化。')

h('4 出生布局与初始状态')
table(['追捕机','期望出生坐标 x y z 米','初始状态'],[
    ['uav_0','100  350  180','速度1米每秒  朝目标'],
    ['uav_1','100  450  200','速度1米每秒  朝目标'],
    ['uav_2','100  550  220','速度1米每秒  朝目标'],
    ['uav_3','100  650  200','速度1米每秒  朝目标'],
], [1.0,3.7,2.0])
p('追捕机在城市西侧纵向集结，逃逸机在东侧出生盒内采样：x为650至750米，y为350至650米，z为150至250米。每次重置会更换目标位置和水平航向，但建筑地图不变。逃逸机初速度为当前级巡航速度、水平随机方向，初始垂直速度为0；双方初始加速度均为0。')
p('实际追捕机坐标经过建筑安全修正，不能把上表期望值无条件当成每级每次重置的最终值。安全净空为机体半径0.5米加保守初步可达距离4米加额外安全余量1.5米，合计6米；修正搜索分辨率1米，边界预留30米，出生之间保留12米间隔。配置中的该出生间隔仍保留，但不产生训练中的机间碰撞惩罚。')
p('目标出生最多拒绝重采样200次，只接受位于出生盒内且不在建筑内的点。安全修正后还核验追捕机到目标的最小真实距离至少450米；无法满足时会报错，不使用越出出生盒的静默兜底。种子1记录的各级初始距离约596至631米，是一次记录示例，不代表所有回合都固定同一初距。')
p('surround_spawn为false，因此课程表遗留的spawn_radius为100至200米、spawn_vertical_radius为30至70米不表示008采用环绕目标出生。它们属于旧模式配置保留值，应以本轮明确的pursuer_spawn_positions和target_spawn_box为准。')
picture('baseline_exp008_spawns_2026-10-13.png','图1 五级固定城市及种子1的真实出生位置  建筑安全修正后的位置以原记录为准')

h('5 三维运动学和时间离散')
h('5 1 实际参数',2)
table(['参数','追捕机','逃逸机'],[
    ['初始速度','1米每秒  向目标','当前级巡航速度  随机水平航向'],
    ['最大速度范数','11米每秒','各级3  5  7  8  9米每秒'],
    ['最大加速度范数','3米每二次方秒','1米每二次方秒  五级相同'],
    ['最大jerk范数','6米每三次方秒','8米每三次方秒'],
    ['加速度滞后时间常数','0.5秒','0.5秒'],
    ['期望速度响应时间','直接输出加速度指令','1秒'],
    ['控制指令更新','每1秒Actor更新一次','每0.1秒势场重新计算'],
    ['物理积分子步','0.1秒  每次决策10个子步','0.1秒'],
    ['最低实际速度','允许为0','巡航速度不是实际速度硬下限'],
], [2.2,2.25,2.25])
p('008使用三维平动、加速度指令、加速度滞后和jerk限制。它没有完整四旋翼姿态、推力、质量、重力补偿、电机响应、风场或电池模型。源文件中还有带yaw的辅助函数，但当前三维动作路径调用advance_translation，不以yaw角速度控制该基线。')
h('5 2 每个物理子步的计算',2)
p('设p、v、a分别为子步开始时的位置、速度和实际加速度，u为期望加速度，h为子步时长。clip_N表示限制三维向量范数，而不是逐轴分别限制。追捕机的三个Actor动作分量首先裁剪到负1至1，再乘3并限制整体范数到3。')
eq(r'j_k=\operatorname{clip}_N((u_k-a_k)/\tau,J_{\max})',
   row(S('j','k'),mo('='),sub(mi('clip'),mi('N')),par(row(frac(row(S('u','k'),mo('−'),S('a','k')),mi('τ')),mo(','),S('J','max')))))
eq(r'p_{k+1}=p_k+v_kh+\tfrac12 a_kh^2',
   row(S('p','k+1'),mo('='),S('p','k'),mo('+'),S('v','k'),mi('h'),mo('+'),frac(mn(1),mn(2)),S('a','k'),sup(mi('h'),mn(2))))
eq(r'v_{k+1}=\operatorname{clip}_N(v_k+a_kh,V_{\max})',
   row(S('v','k+1'),mo('='),sub(mi('clip'),mi('N')),par(row(S('v','k'),mo('+'),S('a','k'),mi('h'),mo(','),S('V','max')))))
eq(r'a_{k+1}=\operatorname{clip}_N(a_k+j_kh,A_{\max})',
   row(S('a','k+1'),mo('='),sub(mi('clip'),mi('N')),par(row(S('a','k'),mo('+'),S('j','k'),mi('h'),mo(','),S('A','max')))))
p('当前τ为0.5秒。位置和速度使用子步开始时的实际加速度；经jerk更新的新加速度从下一子步生效。这是显式离散计算，不是瞬时直接把速度改成Actor动作，也不是加速度滞后的闭式指数解。jerk上限8不能使逃逸机突破1米每二次方秒的加速度限制。')
p('子步数量为ceil(decision_dt除以physics_dt)，实际h为decision_dt除以子步数量。008为10个0.1秒子步。追捕机在这10个子步保持同一指令，但实际加速度持续演化；逃逸机势场每个子步重新读取当前位置、速度和追捕机位置，重新算指令。')
h('5 3 几何接触与数值保护',2)
p('建筑检测覆盖整个位置运动线段，并按机体半径扩大柱体，避免仅看终点而穿过薄障碍。发生建筑进入时，位置截在首次接触前约1毫米；边界把水平位置裁到0.5至999.5米，把高度裁到50至350米。裁剪改变位移时，速度按实际位移除以子步时间重算，加速度状态仍按原动态更新。')
p('裁剪让机器人和目标不持续穿进实体，但接触已被记录并扣分。评估中继续运行的碰撞轨迹可能沿墙滑动或贴边；这属于模拟器保护下的碰撞行为，不能解释为物理上无损、安全的避障。机间接触关闭后，两架追捕机能够重叠或穿过彼此，此设置也限制了真实多机部署解释。')

h('6 逃逸机人工势场与主动逃逸能力')
h('6 1 感知和方向选择',2)
p('逃逸机没有神经网络策略。它使用追捕机排斥、建筑排斥、边界排斥及靠近边缘时的回中偏置构造期望速度。任意追捕机距离小于60米时，目标进入受威胁模式并将期望速度切到本级最大值；否则使用巡航值。在没有明显排斥时保留原速度方向。')
eq(r'f_i=\frac{40}{d_i+5}\frac{p_E-p_i}{d_i}\quad (0.1<d_i<60)',
   row(S('f','i'),mo('='),frac(mn(40),row(S('d','i'),mo('+'),mn(5))),frac(row(S('p','E'),mo('−'),S('p','i')),S('d','i')),mo(','),mn('0.1'),mo('<'),S('d','i'),mo('<'),mn(60)))
table(['参数','当前值','用途'],[
    ['evader_sense','60米','追捕机感知和受威胁速度切换'],
    ['evader_rep_coef','40','追捕机排斥项系数'],
    ['evader_obstacle_range','20米','建筑水平避让影响范围'],
    ['evader_obstacle_coef','40','建筑排斥系数'],
    ['evader_boundary_range','100米','六个空域边界的排斥范围'],
    ['evader_boundary_coef','50','边界排斥系数'],
    ['evader_center_pull','1.5','靠边时回中偏置系数'],
    ['evader_center_xy_range','150米','水平距边界小于该值开始回中'],
    ['evader_center_z_range','75米','垂直距边界小于该值开始回中'],
    ['evader_blend_center','true','将排斥或巡航方向与回中偏置合成'],
], [2.8,1.0,2.9])
p('边界排斥随距边界的1.5次幂倒数增加。建筑排斥依据目标到建筑平面矩形的最近水平距离，在高度不超过楼顶加机体半径加20米时启用，方向主要是水平绕行。若水平投影已经进入建筑附近的极小距离区，普通执行存在随机方向兜底。逃逸机动作之后同样做扫掠检测和位置裁剪，所以空间合法性由势场与几何保护共同维持，不能单靠人工势场宣称从不碰墙。')
h('6 2 回中策略与速度响应',2)
p('回中点是(500,500,200)米。偏置按各轴空域半宽归一化，水平为500米、垂直为150米；只在接近水平150米或垂直75米边缘带时增强，在内部区域为0。开启blend_center后，有排斥则用排斥加回中偏置，无排斥则用巡航惯性方向加回中偏置，再归一化成期望速度。它不是每时刻都强制飞到中心，也没有越过地图边缘就逃逸成功的目标。')
eq(r'u_E=\operatorname{clip}_N((v_{\mathrm{desired}}-v_E)/1,A_E),\quad A_E=1',
   row(S('u','E'),mo('='),sub(mi('clip'),mi('N')),par(row(frac(row(S('v','desired'),mo('−'),S('v','E')),mn(1)),mo(','),S('A','E'))),mo(','),S('A','E'),mo('='),mn(1)))
p('期望方向突然变化并不意味着实际速度瞬时反向。第5级从同向3米每秒加速到9米每秒，理想也至少需要6秒，再叠加滞后；若从反向巡航变为正向逃逸，速度矢量变化可达12米每秒，至少需要约12秒。感知仅60米，特别是第1级50米保持半径，预警距离很短。因此视频中逃逸机可能在捕获前来不及形成明显高速逃跑，并不意味着势场没有被调用。多机排斥向量还可能相互抵消。')
p('增加排斥系数通常改变与墙面、建筑、回中项的相对权重；最终方向会被归一化，实际加速度仍被限制到1。因此单独把系数变大，不能等比例增强物理逃逸加速度。建议后续独立对照感知60、150、200米，优先试200米；该建议尚未实施。约190米的最不利预警粗估来自15米捕获半径、约14米每秒初始闭合速度和约12.5秒转向响应，不能当作保证安全的严格界。')

h('7 策略能够读取哪些环境信息')
p('每架追捕机原始观测79维，额外拼接4维one hot身份向量后进入共享Actor，共83维。信息来自模拟器真值，观测没有噪声、通信延迟、遮挡或目标丢失；逃逸机的60米感知限制并不限制追捕机获得目标状态。')
table(['观测内容','维数','归一化及作用'],[
    ['自身位置  速度  加速度  保持进度','11','位置3  速率1  速度3  加速度3  保持1'],
    ['三架队友相对位置与相对速度','18','每架6维  速度差除以22'],
    ['六个边界距离','6','水平除以1000  垂直除以300'],
    ['目标相对位置  速率  速度  加速度','10','目标速度按当前级最大速度归一化'],
    ['当前课程上下文','4','等级  建筑数量  速度比例  捕获半径'],
    ['最近五个建筑','30','每栋6维  边界相对坐标  顶部相对高度  有效标记'],
    ['原始观测合计','79','11加18加6加10加4加30'],
    ['身份向量','4','共享网络区分uav_0至uav_3'],
], [2.9,.5,3.3])
p('自身水平位置映射到负1至1，高度按50至350米映射；速度和加速度按各自上限归一化。建筑按到三维柱体的最近距离挑选五栋，特征是xmin、xmax、ymin、ymax相对自身位置，楼顶相对高度和有效标记。只输入最近五栋并不等于预先给出全图规划路径。')
p('没有启用目标未来预测观测、图网络或障碍注意力，也没有RNN记忆。环境提供53维集中状态，但本轮MATD3 Critic实际使用联合观测和联合动作，不直接读取该53维state。精确状态与队友信息降低了感知难度，但缺少远处建筑和历史信息也意味着不能无条件把Actor输入视为完整真实世界状态。')

h('8 实际奖励函数')
h('8 1 生效项和整体尺度',2)
p('reward_mode为baseline。每架追捕机的奖励是接近势函数、建筑预警、边界预警、真实碰撞、团队捕获和全队失败的和。所有下文数值已经包含baseline_reward_scale为0.1的缩放，避免把配置中的300误写成实际300分。')
eq(r'r_i=r_{\mathrm{progress},i}+r_{\mathrm{obstacle},i}+r_{\mathrm{boundary},i}+r_{\mathrm{collision},i}+r_{\mathrm{terminal},i}+r_{\mathrm{failure},i}',
   row(S('r','i'),mo('='),S('r','progress,i'),mo('+'),S('r','obstacle,i'),mo('+'),S('r','boundary,i'),mo('+'),S('r','collision,i'),mo('+'),S('r','terminal,i'),mo('+'),S('r','failure,i')))
table(['奖励项','当前实现','触发范围'],[
    ['r_distance','权重0  不贡献总奖励','日志保留0列'],
    ['r_progress','折扣势函数差  权重1  距离尺度10','各机独立接近反馈'],
    ['r_obstacle','负0.5乘预警比例平方乘实际秒数','考虑速度和制动后的建筑净空'],
    ['r_boundary','负0.5乘预警比例平方乘实际秒数','考虑速度和制动后的六面边界净空'],
    ['r_teammate','权重0  不贡献总奖励','不考虑机间避障'],
    ['r_collision','建筑首次负8  边界首次负10  持续每类负2每秒','发生接触的追捕机'],
    ['r_terminal','每架追捕机加30','捕获且本步没有终止型碰撞失败'],
    ['r_failure','每架追捕机负30','训练中的建筑或边界碰撞提前终止'],
], [1.2,3.6,1.9])
p('覆盖、准备度、闭合度、速度匹配、保持塑形和每步时间成本均不参与baseline总奖励；保持2秒仍用于成功判定。配置中遗留的w_prep、w_cov等旧模式值不是本轮有效奖励。008没有单独动作幅度、能耗或平滑度奖励，平滑性来自动态滞后与jerk约束。')
h('8 2 接近势函数及终止处理',2)
eq(r'\Phi_i(s)=-d_i(s)/10,\qquad r_{\mathrm{progress},i}=0.1[0.99\Phi_i(s^{\prime})-\Phi_i(s)]',
   row(sub(mi('Φ'),mi('i')),par(mi('s')),mo('='),mo('−'),frac(sub(mi('d'),mi('i')),mn(10)),mo(','),S('r','progress,i'),mo('='),mn('0.1'),mo('['),mn('0.99'),sub(mi('Φ'),mi('i')),par(mi('s′')),mo('−'),sub(mi('Φ'),mi('i')),par(mi('s')),mo(']')))
p('非终止步可写成0.01乘以“本步前距离减0.99乘本步后距离”。从600米接近到590米，该项为0.159；600米保持不动时为0.060；退到610米时为负0.039。这一项不是不动严格为0的简单距离差，小幅退远也可能由于折扣而得到正值。')
p('捕获或真正碰撞终止时，下一状态势值取0；超时截断仍保留下一势值，用于价值自举。终止零势值会使最后一步得到正的势函数项，即使是碰撞失败，也不会自动成为负progress；失败由单独的负30团队惩罚表达。仅看最后一步progress的正负不足以判断任务完成。')
p('折扣势函数塑形在满足原MDP与终止处理等条件时有保持最优策略的理论依据，但并不保证这个共享参数、多智能体、神经网络近似和有限训练预算的问题收敛。参考Ng、Harada和Russell的1999年势函数塑形论文；本轮公式和缩放以源码为准。')
blocks.append(('web','Policy invariance under reward transformations','https://people.eecs.berkeley.edu/~russell/papers/icml99-shaping.pdf'))
h('8 3 带制动余量的稠密避障',2)
p('本轮解析后的运行配置启用了baseline_braking_safety和baseline_braking_deficit_penalty，反应时间显式为1.5秒。接近速度v正部取朝向建筑最近点或相应边界面的速度投影，再从几何净空中扣掉反应距离和理想制动距离。')
eq(r'D(v^+)=1.5v^++(v^+)^2/(2\times3),\qquad c=d-\rho-D(v^+)',
   row(mi('D'),par(sup(mi('v'),mo('+'))),mo('='),mn('1.5'),sup(mi('v'),mo('+')),mo('+'),frac(sup(par(sup(mi('v'),mo('+'))),mn(2)),row(mn(2),mo('×'),mn(3))),mo(','),mi('c'),mo('='),mi('d'),mo('−'),mi('ρ'),mo('−'),mi('D'),par(sup(mi('v'),mo('+')))))
eq(r'r_{\mathrm{avoid}}=-0.5\left[\max(0,1-c/50)\right]^2\Delta t',
   row(S('r','avoid'),mo('='),mo('−'),mn('0.5'),sup(row(mo('['),mi('max'),par(row(mn(0),mo(','),mn(1),mo('−'),frac(mi('c'),mn(50)))),mo(']')),mn(2)),mi('Δt')))
p('建筑和边界分别取各自最危险的有效净空c，预警范围均为50米。启用deficit后，c可以小于0，预警比例可以大于1，惩罚没有被限制在每秒负0.5以内。以朝向正面障碍11米每秒为例，反应和制动距离约36.67米，开始产生预警的实际净空可达约86.67米。该计算是局部制动近似，未显式求解最优转弯或绕楼路线。')
p('建筑净空实现按最近柱体点投影，顶部低于最低可飞高度的建筑也仍存在于源数据中。因此它可能对实际可飞高度下方楼体产生预警偏置；这是后续需要核查的奖励与几何一致性问题，不能只按建筑栋数解释训练难度。')
h('8 4 接触事件和全队失败',2)
p('首次进入建筑接触扣8，首次进入边界接触扣10；持续接触按每架、每类每秒扣2。离开接触后再进入可以再次产生首次事件扣分。同一物理子步同时接触建筑与边界时两项分别累计。训练通常很快因首次真实碰撞结束，评估继续运行可能累积很大的持续接触惩罚。')
p('r_failure在实验003加入；既有源快照中实验001没有该项，实验003已有。003于2026年10月13日01:17北京时间启动。008保留它，用于全队因单机建筑或边界碰撞而提前失败时，给每架追捕机负30；肇事机还会收到自己的r_collision。它不是本轮临时新增的超时罚分，400步超时不触发该项。')
p('评估关闭碰撞提前终止，因此评估中的r_failure通常为0；若碰撞后后来捕获，终止奖励分支也可能加30，但安全成功仍为0。训练和全回合评估的累计回报受终止方式影响，并非完全同一口径。比较策略首先看安全成功、接触率和捕获时间，不能只比较两者mean_return。')

h('9 MATD3网络与奖励影响学习的路径')
h('9 1 算法复用及网络结构',2)
p('当前算法复用XuanCe已有MATD3的Agent、Policy和Learner，并沿用其多智能体训练和回放基础；没有把用户的MADDPG全部重新写成一套独立算法。MATD3策略类继承MADDPG策略类，Agent沿用IDDPG相关基础。在此基础上调整了本环境、奖励、课程、日志以及TD3目标平滑和Actor使用第一张Critic的配置。')
table(['网络','输入到输出','当前处理'],[
    ['共享Actor','83 → 512 → 512 → 3','LeakyReLU隐藏层  Tanh动作'],
    ['共享Critic A','332 → 512 → 512 → 1','联合观测 联合动作 当前身份'],
    ['共享Critic B','332 → 512 → 512 → 1','独立Q参数  与A同时拟合'],
    ['目标网络','Actor与两张Critic的目标副本','软更新系数0.001'],
], [1.2,2.6,2.9])
p('Critic输入为四架79维观测的拼接316维，加四架3维动作12维，加当前代理身份4维，共332维。使用Basic_Identical表征，不在上述MLP之前再增加一个隐藏特征网络。共享Actor结合身份可以学习不同分工，但共享参数本身并不保证学会包围或协作。')
h('9 2 Critic目标与Actor梯度',2)
p('经验回放保存联合观测、实际探索动作、奖励、下一观测和终止标记。环境计算由普通数值代码完成，没有沿物理仿真把奖励直接反向传播到Actor。奖励首先作为Critic的监督目标之一，再通过Critic关于动作的梯度影响Actor。')
eq(r'a^{\prime}=\operatorname{clip}(\mu^{\prime}(o^{\prime})+\operatorname{clip}(\epsilon,-0.5,0.5),-1,1),\quad \epsilon\sim N(0,0.2^2)',
   row(mi('a′'),mo('='),mi('clip'),par(row(mi('μ′'),par(mi('o′')),mo('+'),mi('clip'),par(row(mi('ε'),mo(','),mo('−'),mn('0.5'),mo(','),mn('0.5'))),mo(','),mo('−'),mn(1),mo(','),mn(1))),mo(','),mi('ε'),mo('∼'),mi('N'),par(row(mn(0),mo(','),sup(mn('0.2'),mn(2))))))
eq(r'y_i=r_i+\gamma(1-done_i)\min(Q_A^{\prime}(O^{\prime},a^{\prime},i),Q_B^{\prime}(O^{\prime},a^{\prime},i))',
   row(S('y','i'),mo('='),S('r','i'),mo('+'),mi('γ'),par(row(mn(1),mo('−'),S('done','i'))),mi('min'),par(row(sup(S('Q','A'),mo('′')),par(row(mi('O′'),mo(','),mi('a′'),mo(','),mi('i'))),mo(','),sup(S('Q','B'),mo('′')),par(row(mi('O′'),mo(','),mi('a′'),mo(','),mi('i')))))))
eq(r'L_C=\frac{1}{4B}\sum_{b,i}\left[(Q_{A,i}-y_i)^2+(Q_{B,i}-y_i)^2\right]',
   row(S('L','C'),mo('='),frac(mn(1),row(mn(4),mi('B'))),sub(mo('∑'),mi('b,i')),mo('['),sup(par(row(S('Q','A,i'),mo('−'),S('y','i'))),mn(2)),mo('+'),sup(par(row(S('Q','B,i'),mo('−'),S('y','i'))),mn(2)),mo(']')))
eq(r'L_A=-\frac{1}{4B}\sum_{b,i}Q_{A,i}(O_b,A_{\theta}(O_b))',
   row(S('L','A'),mo('='),mo('−'),frac(mn(1),row(mn(4),mi('B'))),sub(mo('∑'),mi('b,i')),S('Q','A,i'),par(row(S('O','b'),mo(','),sub(mi('A'),mi('θ')),par(S('O','b'))))))
p('B为批量512，i遍历四个代理。目标取两张目标Critic的较小值，Critic误差是两张Q的平方误差之和再求均值；Actor的配置td3_actor_use_first_critic为true，因此优化第一张Critic A对当前策略联合动作的Q。Actor每两次Critic更新一次，降低Actor追随快速变化Q估计的频率。')
p('由于共享Actor对四架联合动作都重新生成，链式求导会沿Q对各架动作的偏导传播到同一套Actor参数；并不是只更新当前索引i对应一架的参数。可以把链条理解为：奖励改变TD目标，TD误差改变Critic，Critic的动作梯度告诉Actor哪些指令有较高长期回报。这里没有直接对奖励函数解析求导，也不能用reward的某个正值断言Actor梯度必然朝目标方向。')
p('双Critic取小值、目标动作平滑和延迟策略更新来自TD3方法，用于减轻价值估计误差；它们并不保证本任务不震荡或在给定预算内收敛。Actor loss可以为负；训练课程改变会改变数据分布、奖励规模和Q值。判断训练应结合安全成功、碰撞类别、动作饱和和Q变化，而不能要求loss_actor收敛到0。')
blocks.append(('web','Addressing Function Approximation Error in Actor Critic Methods','https://arxiv.org/abs/1802.09477'))

h('10 训练超参数与初始化')
table(['设置','实际值','说明'],[
    ['Actor学习率','0.00003','Adam'],
    ['Critic学习率','0.00015','两张Q同时更新'],
    ['Adam epsilon','0.00001','沿用Learner优化器实现'],
    ['折扣gamma','0.99','每1秒决策的折扣  也用于progress'],
    ['目标软更新tau','0.001','随延迟Actor更新执行'],
    ['学习率末端比例','1.0','本轮不衰减学习率'],
    ['batch size','512','联合多智能体转移'],
    ['buffer size','200,000','不是每级独立20万容量'],
    ['学习开始','4,096环境转移','初始化权重也重新收集经验'],
    ['训练频率','1','每次向量交互后一个Learner迭代'],
    ['Actor更新间隔','2次Critic更新','实际计数为Critic的一半'],
    ['目标策略噪声','标准差0.2  裁剪至正负0.5','加到目标动作后再裁剪动作范围'],
    ['训练探索噪声','0.1降到0.03','前300,000转移衰减  后续保持0.03'],
    ['梯度范数裁剪','0.5','Actor和Critic均启用'],
    ['参数共享','true','四机同一Actor和两张共享Critic'],
    ['图模块和障碍GAT','false','纯MLP  不使用RNN'],
    ['训练种子','1','本轮只有一个训练种子'],
    ['正式评估频率','每50,000环境转移','每次100回合  8并行'],
    ['跨级第5级诊断','每250,000环境转移','phase为final_task  不表示已训练第5级'],
], [2.15,2.05,2.5])
p('解析配置中的通用learning_rate为0.0004是保留字段，实际Actor与Critic使用各自学习率。sigma为0.1也保留在继承配置中，实际探索记录应以noise_scale和start_noise、end_noise、noise_decay_steps为准。为避免误读，完整config.json随原始记录保留，复现时必须同时看继承基础配置与本轮YAML合并后的有效值。')
p('初始化来源是007冻结最终权重，校验了模型SHA256并完整加载策略相关张量，包括目标网络；经验回放和优化器新建。因此008是在新出生布局和新奖励设置下的迁移续训，不是007回放和优化器状态的逐步恢复，也不是随机权重基线。旧007在原环绕出生任务的94%安全成绩不能代替它在008新布局上的成绩。')

h('11 五级课程与升级方式')
table(['显示级别','内部level','建筑数量','巡航及受威胁速度','保持捕获半径'],[
    ['1','0','20','1和3米每秒','50米'],
    ['2','1','40','2和5米每秒','35米'],
    ['3','2','60','3和7米每秒','25米'],
    ['4','3','80','3和8米每秒','20米'],
    ['5','4','100','3和9米每秒','15米'],
], [.7,.95,.9,2.65,1.5])
p('各级同时改变建筑数量、目标巡航及最大期望速度、连续保持捕获半径。目标加速度仍为1，直接拦截半径仍为5米，连续保持仍为2秒。课程没有再给覆盖设置门槛，也没有自动按预算强行升级。')
p('升级由PerformanceCurriculum处理：当前级至少经历100,000环境转移后，连续两次不同评估步的有效100回合评估都满足安全成功率至少80%且总碰撞率至多10%，才进入下一级。有一次未通过则连续计数重置。当前max_level为内部4；到达第5级后不会再升级，必须另外据成绩判断是否达标。')
p('这一门槛不是80%捕获率。第1级135万步时安全80%但碰撞20%，仍不能通过；140万和145万时分别97%与100%安全，且碰撞3%与0%，才于145万升级。另一个旧的环境update_curriculum函数仍存在，但本轮benchmark使用的是上述性能课程，不使用旧函数的降级规则。')
p('升级消息在各并行环境下一次reset时应用。尚未结束的旧级回合继续保留原级设置，完成后再转新级；因此完整回合按自身level统计的转移数可以略超过全局课程区间，不能强行把CSV改成整齐的区间预算。当前训练只采样当前课程，没有定期混入旧级回合，升级评估也默认不复测全部旧级。')

h('12 实际训练量与阶段效果')
h('12 1 回合 转移和轮数',2)
p('这里的一个完整回合是环境从reset到捕获、碰撞提前终止或超时的一次轨迹。一次环境转移是某个并行环境执行一个决策；16环境同时执行一步通常是16转移。强化学习没有监督学习那种固定“每轮遍历全数据”的epoch，不能把50,000步评估周期或一次优化批量叫训练一轮。')
intervals = [1450000,600000,150000,550000,250000]
stage_stats = []
for i in range(5):
    rr = [r for r in episodes if int(r['level']) == i]
    st = {'level':i+1,'episodes':len(rr),'completed_transitions':sum(int(r['episode_steps']) for r in rr),
          'safe_success_count':sum(int(r['safe_success']) for r in rr)}
    stage_stats.append(st)
table(['级别','全局课程区间转移','完整回合数','完整回合转移','训练安全成功数'],[
    [str(s['level']),f'{intervals[i]:,}',f"{s['episodes']:,}",f"{s['completed_transitions']:,}",f"{s['safe_success_count']:,}"]
    for i,s in enumerate(stage_stats)], [.55,1.7,1.3,1.6,1.55])
complete = sum(s['completed_transitions'] for s in stage_stats)
p(f'合计{len(episodes):,}个完整回合，完整回合转移{complete:,}，训练终点还有{3000000-complete:,}个转移属于未完成回合，未写入完整回合CSV。各级累计训练安全成功比例混合早期探索失败与后期能力，不等同于无探索正式评估。总优化计数187,244次Critic、93,622次Actor以真实training_updates.csv为准。')
h('12 2 正式评估的历史最好和最后一次',2)
table(['级别','最好步数','最好捕获 安全 碰撞','末次步数','末次捕获 安全 碰撞'],[
    [str(i+1),f"{summary['best_by_stage'][str(i)]['step']/10000:.0f}万",
     ' / '.join(pct(summary['best_by_stage'][str(i)][k]) for k in ['success_rate','safe_success_rate','collision_rate']),
     f"{int([r for r in current if int(r['level'])==i][-1]['step'])/10000:.0f}万",
     ' / '.join(pct([r for r in current if int(r['level'])==i][-1][k]) for k in ['success_rate','safe_success_rate','collision_rate'])]
    for i in range(5)], [.55,.85,2.15,.85,2.3])
p('四次课程升级发生在145万、205万、220万和275万步，分别进入第2至5级。第1级前125万步正式评估安全成功一直为0，多数是边界接触；130万出现18%安全成功，之后达到升级门槛。早期诊断发现动作分量接近正负1、策略朝边界饱和，这是单条轨迹诊断，不能替代成功率证据。')
p('第2级开始时建筑接触率100%、安全成功0%；后续曾出现明显反复，例如190万安全95%，195万降到69%，最终200万和205万连续通过。第3级只训练约15万全局转移便通过。第4级初期安全25%左右，270万与275万分别96%和91%，达到连续两次条件。课程数量为5符合此前4至6级的设计范围，但学习难度并非逐级等距。')
p('第5级在280万步安全1%，285万提升到59%，290万最高74%，295万降到56%，300万为65%。捕获92%的最好记录与安全74%是同次结果，不能把捕获率当安全率。第5级最后25万转移表现仍不稳定，也未达到门槛；正常完成预算只表示运行成功结束。')
picture('baseline_exp008_final_curriculum_2026-10-13.png','图2 全部300万步正式评估  实线为安全成功 虚线为捕获  竖线表示真实课程升级')

h('13 最终模型复评与旧级能力保留')
h('13 1 同一最终权重的五级复评',2)
table(['级别','捕获率','安全成功','总碰撞','建筑','边界','平均决策步'],[
    [str(i+1)]+[pct(r['metrics'][k]) for k in ['success_rate','safe_success_rate','collision_rate','building_collision_rate','boundary_collision_rate']]+[f"{r['metrics']['mean_episode_steps']:.2f}"]
    for i,r in enumerate(finals)], [.5,.85,1.0,.9,.75,.75,1.95])
p('上述每行是种子60001的100回合复评，无探索，使用008原始1秒决策和0.1秒物理配置。平均决策步包含失败、超时和成功混合回合，不能当作成功捕获时间均值。因为一回合最后决策可能在捕获分数时刻结束，平均决策步也不是精确物理秒数均值。')
p('第1至3级虽然捕获都是100%，但只有38%、4%、6%回合全程无接触。最终策略常能最终追上，却失去旧级安全路径控制。第4级安全68%，第5级57%，均未达到既定标准；因此最终单一模型不能宣布五级课程毕业。')
h('13 2 与007初始化权重比较',2)
table(['级别','007在新布局捕获','007在新布局安全','008最终捕获','008最终安全'],[
    [str(i+1),pct(initials[i]['metrics']['success_rate']),pct(initials[i]['metrics']['safe_success_rate']),pct(finals[i]['metrics']['success_rate']),pct(finals[i]['metrics']['safe_success_rate'])]
    for i in range(5)], [.55,1.6,1.6,1.5,1.45])
p('初始化007权重在相同008新布局和60001种子下，五级安全均为0、碰撞均为100%。008最终安全成绩已改善，证明续训确实改变了新任务表现；这张对照表不评价007原任务，不能和其旧布局94%直接混为一次下降。')
h('13 3 同种子对照支持旧级能力遗忘',2)
table(['级别','阶段最好模型步数','阶段最好安全','最终模型安全','阶段最好碰撞'],[
    [str(i+1),f"{summary['best_by_stage'][str(i)]['step']/10000:.0f}万",pct(retained[i]['metrics']['safe_success_rate']),pct(finals[i]['metrics']['safe_success_rate']),pct(retained[i]['metrics']['collision_rate'])]
    for i in range(3)], [.55,1.9,1.4,1.35,1.5])
p('三个阶段模型在同一种子上仍分别达到100%、99%、100%安全，而最终模型明显下降。这使“只因换了出生种子而变差”不能充分解释退化，支持旧级安全能力遗忘。阶段模型在CPU、最终模型在CUDA执行，均为同一float32策略和无探索流程；设备并非严格同一，但差异幅度很大。该检查不是多训练种子统计检验，也没有证明某一个超参数就是唯一原因。')
p('值得优先检验的机制是：每次升级后只训练新级；20万回放容量会被新级经验覆盖；第5级训练25万转移，末尾旧级经验占比可能很低；共享权重在新速度和密度下调整而未接受旧级保留约束。本轮未保存回放分级占比，不能把这一机制当作已经完成的因果消融结论。')
picture('baseline_exp008_final_retention_2026-10-13.png','图3 新布局迁移收益与旧级保留对照  全部使用60001种子每级100回合')

h('14 奖励 逃逸行为与优化曲线')
p('为观察各级训练末尾的策略状态，下面使用每级最后100个完整训练回合。奖励为四机平均的回合累计分项，已乘0.1；不是每步奖励，也不是正式评估回报。四机独立稠密项求平均，团队捕获和失败项同值。')
reward_keys = ['r_progress','r_obstacle','r_boundary','r_collision','r_terminal','r_failure']
last100 = [[r for r in episodes if int(r['level']) == i][-100:] for i in range(5)]
means = [{k:sum(float(r[k]) for r in rr)/len(rr) for k in reward_keys+['target_speed_mean','target_sensed_fraction']} for rr in last100]
table(['级别','progress','建筑预警','边界预警','接触','捕获','失败'],[
    [str(i+1)]+[f'{m[k]:.3f}' for k in reward_keys] for i,m in enumerate(means)], [.55,1.15,1.05,1.05,.9,1.0,1.0])
table(['级别','目标平均实际速度 米每秒','目标感知比例'],[
    [str(i+1),f"{m['target_speed_mean']:.3f}",pct(m['target_sensed_fraction'])] for i,m in enumerate(means)], [.65,3.5,2.55])
p('目标感知比例是回合决策末采样中“至少一架追捕机进入60米”的比例，再对回合求均值，不是0.1秒连续感知时长占比。第5级实际目标平均速度仍约3.21米每秒，接近巡航值3，而非每时刻9米每秒。这与晚感知、加速度限制和转向过程相符；仅凭最大速度配置9米每秒不足以认为训练已经面对持续高速积极逃逸。')
p('最后优化窗口Actor loss约负5.432、Critic loss约11.086、predictQ约4.368。predictQ在本轮Learner日志中是回放动作对应两张Q的较小值的均值；Actor loss对应当前策略联合动作在Critic A下的负均值，两者输入动作不同，不能要求loss_actor严格等于负predictQ。旧界面称pre_Q时，应先核对其来源；008实际字段为predictQ。')
picture('baseline_exp008_final_rewards_and_losses_2026-10-13.png','图4 完整回合奖励窗口与实际优化日志  曲线有限或下降都不能单独证明收敛')

h('15 五级视频和样例口径')
p('五级均已渲染真实无碰撞成功样例。每个视频使用该级历史最好权重，并在原环境逐动作重放，核验轨迹一致、模型哈希一致、建筑和边界接触均为0。它们说明该级保存模型存在成功轨迹，不代表最终模型表现，也不代替100回合成功率。全部按2倍速度播放。')
table(['级别','模型训练步数','筛选种子','捕获时刻 秒','捕获方式'],[
    ['1','145万','62001','64.30','单机连续保持'],
    ['2','200万','63001','70.27','单机连续保持'],
    ['3','210万','63001','76.24','单机连续保持'],
    ['4','270万','63001','74.77','单机连续保持'],
    ['5','290万','63001','81.20','单机连续保持'],
], [.55,1.55,1.15,1.5,1.95])
for i,(relative,v) in enumerate(videos,1):
    link(f'第{i}级成功视频', relative+'/success.mp4')
p('62001和63001已经用于视频筛选，60001已经用于本轮最终复评，不应在下一实验中再称作未使用盲测种子。第1级视频复用原008交付，不重新渲染；第2至5级在新的唯一目录保存，未混用007视频。各目录同时保存manifest、模型、动作轨迹和关键帧。')

h('16 日志字段与TensorBoard查看')
p('TensorBoard服务器读取008 seed根事件目录：outputs/logs/matd3/baseline_exp008/seed_1_2026_1013_183201。通过现有SSH转发访问 http://127.0.0.1:6006/。日志总量较大，本次交付包含原始CSV和复评JSON，但没有声称完整TensorBoard事件已全部复制到本机。')
table(['日志组或文件','读取方法'],[
    ['Train Rewards 各机和team','每步均值  回合累计分项  当前级最近100回合均值'],
    ['训练回报与结果','四机各自总回报  捕获  安全成功  碰撞  超时'],
    ['运动与距离','初始和最终距离  最近距离  速度  加速度  目标速度和感知比例'],
    ['Train Learning 与Learner','优化窗口均值  实际Critic和Actor计数  loss和Q'],
    ['Test Results','无探索正式评估  必须区分current与final_task'],
    ['training_episodes.csv','只包含完整回合  level为零基  分项为每机平均累计'],
    ['training_updates.csv','优化窗口计数与均值  updates和actor_updates为累计真实更新'],
    ['curriculum_eval.csv','100回合正式评估  phase为current或第5级跨级诊断'],
    ['training_status.json','运行中的即时快照  完成后可能稍早于真实终点'],
    ['baseline_summary.json','最终总步数  课程切换记录  历史最好成绩'],
], [2.25,4.45])
p('此前TensorBoard错误读取006的uav_3/r_failure/env子目录，导致界面只看见很少分项。examples/run_tensorboard.py已修复：seed根有事件文件时直接读取根，不继续选择分项子文件夹。日志改为根事件文件的平面标量，覆盖所有无人机，避免一个分项一个子运行造成难以查找。共享参数优化主要出现在uav_0标记下，不表示其他三架没有更新。')
p('008启动前47项测试通过；GPU集成测试2048转移，真实Critic更新1008次、Actor504次，338个根标量标签均有限，且没有分项子目录。这些验证说明代码和日志路径可执行，并不替代长期实验成绩，也不构成loss收敛结论。')

h('17 本轮已经实施的改动')
p('奖励方面：关闭r_distance，只保留折扣progress塑形；关闭r_teammate和旧分离项；保留带制动余量的建筑和边界稠密避障、真实接触、捕获和全队失败惩罚。覆盖奖励暂不参与训练，成功条件仍沿用单机拦截或连续保持。')
p('碰撞方面：关闭机间碰撞开关与机间终止，并修复关闭后仍累计冲突、污染安全成功率的问题；建筑和边界仍做完整检测。正式评估关闭碰撞提前终止且累计接触，避免“撞一下就结束”掩盖策略随后持续撞墙。')
p('环境方面：固定五张均匀嵌套城市，按50%、40%、10%建筑类型构成；保持20至100栋的五级数量、速度与捕获半径课程。008采用西侧四机集结和东侧目标出生盒，最小真实初距450米，保留出生建筑安全修正。')
p('观测和日志方面：保留目标和队友运动信息、最近五栋建筑与课程上下文；新增所有无人机回报、奖励分项、完整回合结果、出生与最终距离、逃逸速度、感知比例和真实优化计数；修复TensorBoard根目录选择。')
p('实验管理方面：冻结007初始化来源与29项训练源哈希，记录真实退出码；保留008原预算和配置直到正常结束；完成最终权重五级独立种子复评、007新布局对照、旧级阶段模型保留对照，以及五级成功视频。现有007冻结交付和旧视频未被覆盖。')

h('18 环境能学到什么追捕策略')
p('从已核实数据看，该设置可以学到根据当前目标位置和速度调整三维加速度、持续接近、保持在捕获半径以及部分绕行或飞越行为。阶段模型在新出生种子上仍有高安全成功，成功视频在原动态环境中重放一致，说明行为来自实际策略闭环执行。')
p('任务定义仍主要奖励任一单机捕获，没有要求四机同时形成包围、切断不同逃逸方向或共同承担拦截角色。四机联合信息与共享成功奖励可以促成协作，但不能用任意单机成功自动证明学到了多机围捕几何。若论文主张协同围捕，需要额外报告角色分工、目标周围角度分布、协同带来的收益和单机对照。')
p('当前实验是固定地图、理想状态反馈、弱主动逃逸目标下的仿真基线。目标加速度低、感知晚，追捕机更快且拥有全程准确目标状态；很多楼体低于出生高度，允许飞越；机间碰撞忽略。这些都是任务已明确采用的简化，需要在论文中写清，而不能将结果直接解释为真实传感器条件、受风四旋翼或所有城市的可靠追捕。')
p('最终模型安全复评低于门槛，加上旧级能力遗忘，说明当前“到达第五级”不足以证明最终策略已稳定学会整个任务。单训练种子、100回合评估以及固定城市也不足以估计跨训练随机性的稳定性或未见布局泛化。后续应逐项增强任务，并使用独立模型对照，保留一个明确可复现的简单基线。')
blocks.append(('web','A Study on Overfitting in Deep Reinforcement Learning','https://arxiv.org/abs/1810.12282'))

h('19 待做事项与下一实验设计')
h('19 1 已采纳的时间设置',2)
p('下一实验采用物理0.1秒、决策0.5秒：Actor控制频率提高到2赫兹，逃逸势场和物理仍为10赫兹，每次决策有5个子步。008作为原始1秒决策对照保留，不覆盖其模型、CSV和归档。尚未启动下一长实验。')
table(['项目','008当前','0.5秒决策的待设值'],[
    ['decision_dt','1.0秒','0.5秒  已采纳'],
    ['physics_dt','0.1秒','0.1秒  保持'],
    ['episode_length','400','800  保持最多400秒'],
    ['capture_hold_required','2秒','2秒  使用真实时间'],
    ['gamma','0.99','约0.994987  若保持每秒折扣一致'],
    ['progress gamma','0.99','与Learner同步  不能漏改环境'],
    ['制动反应时间','显式1.5秒','明确选择1.0秒或保留1.5秒  写入配置'],
], [2.15,1.6,2.95])
p('保持每秒折扣时，新gamma为0.99的平方根，约0.994987437。按实际秒数积分的建筑、边界和持续接触奖励无需机械再除二；事件捕获与失败仍各一次，progress也不能随意额外乘0.5。008解析配置显式继承反应时间1.5秒，仅改decision_dt不一定会触发构造函数的“决策时间加滞后”默认值；必须在下一配置明确记录反应时间。')
p('若为了控制单一因素，首轮可保留反应时间1.5秒；若要按较快控制响应同步修改，可设1.0秒并把它记为第二项改变。哪种设计更合适取决于要比较控制频率本身还是完整响应方案，不能省略这一配置差异。')
h('19 2 公平比较训练预算',2)
p('300万转移乘0.5秒只有原设置一半的名义交互时长；要匹配名义物理交互时长需约600万转移。捕获分数步使真实模拟秒数略不同，最好额外记录实际elapsed秒总和。固定转移比较和固定物理时长比较回答不同问题，不能把两者作为完全相同预算。')
p('若按物理时间匹配，候选同步调整为：最大步800、探索衰减60万、暖启动8192、课程最少20万、评估间隔10万、回放40万。上述均是待评估的预算方案，不是已运行设置，也不自动授权启动600万长实验。学习频率和Actor延迟还会影响每秒梯度更新数，应同时报告真实更新计数。')
h('19 3 优先解决旧级能力遗忘',2)
p('优先级P0：加入旧级复测与保留机制。每次升级、周期评估和最终选模都检查已到达各级；除当前级安全与碰撞门槛外，保留旧级下限。模型选择按多级共同安全表现，不能仅按当前最后级最好分保存一个通用模型。')
p('建议做独立对照：纯当前级采样与当前级70%、已学级30%混合采样；或按级维护有最低保留量的回放。70比30是初始试验建议，尚未实现或证明最优。逐项记录回放级别占比、各级安全率、动作饱和比例和每级奖励，判断保留旧经验是否有效。')
p('对照阶段最好权重可用于判断能否快速恢复旧级安全，但不能以部署时偷偷换best_stage模型替代单一最终策略。若允许按任务切换阶段模型，应明确这是多模型方案，另行评估切换策略。')
h('19 4 提升逃逸能力与检查几何一致性',2)
p('优先级P1：在已确定时间步设置之后，保持其他条件冻结，测试感知60、150、200米；先试200米，核对进入感知后的反应时延、实际最大及平均速度、航向变化、排斥向量是否抵消，以及捕获之前有多少真实主动逃逸时间。必要时再单独研究目标加速度，当前1并未改成其他值。')
p('优先级P1：核查建筑顶部低于z_min时，制动预警使用的最近点是否与扫掠实体一致；核查捕获与碰撞在同一0.1秒子步不同分数时刻发生时，碰撞事件与持续接触罚分是否严格按事件先后截断。当前实现先统计整子步接触再插值捕获，存在需要验证的边界时序风险；这里列为待测，不把它写成已证实错误。')
p('优先级P1：对第5级失败回合分类，分别抽取建筑、边界、超时、捕获但不安全样例，绘制逐轴动作、速度、加速度、有效净空、progress和Q。不要因为loss低而停止检查行为，也不要通过强行升级改写008结果。')
h('19 5 扩展证据与论文实验',2)
p('优先级P2：至少增加多个独立训练种子，并对冻结模型使用独立评估种子；报告安全成功、捕获、各类接触、成功回合捕获时间、全部回合截断情况及均值波动。评估复用当前五图与未见城市布局应分别给出。')
p('优先级P2：做共享团队奖励与单机策略对照，检查四机包围角度与分工；如果最终研究目标确实是严格围捕，应先改变成功定义再设计相应奖励，并保留当前单机拦截或保持作为简单baseline。覆盖暂时不加入这一轮。')
p('优先级P2：逐步增加观测噪声、目标遮挡、通信延迟、风扰动、姿态或底层控制误差，以及真实机间安全约束。每增加一项用独立实验编号保存，避免把多项改变带来的结果误归因给某一个奖励。')
table(['状态','事项','下一交付'],[
    ['已完成','008训练与13组结束后评估','原始记录  哈希  本报告'],
    ['已完成','五级无碰撞成功视频','阶段模型来源与重放核验'],
    ['已采纳 待实施','0.1秒物理加0.5秒决策','新配置  时间尺度与预算核对'],
    ['建议 待实施 P0','旧级复测  混合课程或分级回放','保留能力对照'],
    ['建议 待实施 P1','更远目标感知与接触时序检查','单项对照与失败诊断'],
    ['建议 待实施 P2','多种子 新城市  协同与真实条件','泛化与鲁棒性证据'],
    ['保持暂停','持续监控uav','按需查看  不自动重启训练'],
], [1.65,3.05,2.0])

h('20 复现 文件索引与校验')
h('20 1 运行源码与配置',2)
p('运行任何xuance脚本都必须显式把PYTHONPATH设为独立工作树。服务器默认editable安装可能指向原UAV_Mine，若省略会错误导入旧环境。训练和此次复评使用正确的独立源码。复现前先检查metadata中的源哈希及分支，保留原任务修改。')
code('cd /home/ryy/UAV_Mine_baseline_coverage/xuance-master\nexport PYTHONPATH="$PWD"\n# 以下仅为原实验命令记录 不要重复启动同名008\n/home/ryy/miniconda3/envs/rl_env/bin/python -u examples/run_baseline.py \\\n  --name exp008 --config experiments/baseline_exp008.yaml \\\n  --steps 3000000 --parallels 16 --seed 1 --device cuda:0 \\\n  --start-level 0 --eval-interval 50000 --test-episodes 100')
table(['文件','主要作用'],[
    ['experiments/baseline_exp008.yaml','本轮覆盖配置  有效配置需结合继承基础'],
    ['outputs/results/matd3/baseline_exp008/config.json','完整解析运行参数  含制动开关与反应时间'],
    ['xuance/environment/multi_agent_env/uav_pursuit_coverage_3d.py','观测 出生  势场  捕获 接触 奖励'],
    ['xuance/environment/multi_agent_env/uav_kinematics_3d.py','显式平动积分与范数限制'],
    ['xuance/engine/uav_curriculum.py','连续两次性能门槛升级'],
    ['xuance/torch/learners/multi_agent_rl/matd3_learner.py','目标Q loss Actor梯度与计数'],
    ['tools/generate_uniform_city_maps.py','五张确定性嵌套城市'],
    ['examples/evaluate_baseline_checkpoint.py','冻结模型逐级无探索评估'],
    ['tools/render_baseline_success_videos.py','成功筛选与原环境逐动作重放'],
], [3.3,3.4])
h('20 2 冻结模型和来源',2)
table(['模型','SHA256'],[
    ['007初始化模型',manifest['initial_sha256']],
    ['008最终模型',manifest['final_sha256']],
    ['第1级最好模型',retained[0]['checkpoint_sha256']],
    ['第2级最好模型',retained[1]['checkpoint_sha256']],
    ['第3级最好模型',retained[2]['checkpoint_sha256']],
    ['第4级最好模型',videos[3][1]['checkpoint_sha256']],
    ['第5级最好模型',videos[4][1]['checkpoint_sha256']],
], [1.7,5.0])
p('最终冻结模型为outputs/reviews/exp008_final_v1/final008.pth，来源是训练输出final_train_model.pth；阶段模型来源为best_stage_0至best_stage_4。末级历史最好模型在290万步保存，最终权重在300万步，二者不可混用。哈希在冻结、评估和本机传输后校验。')
h('20 3 本机可打开的原始材料',2)
for label,path in [
    ('完整有效运行配置','outputs/results/matd3/baseline_exp008/config.json'),
    ('本轮冻结训练源快照','outputs/experiments/exp008/source_snapshot.zip'),
    ('完整训练回合CSV','outputs/results/matd3/baseline_exp008/training_episodes.csv'),
    ('完整优化更新CSV','outputs/results/matd3/baseline_exp008/training_updates.csv'),
    ('完整课程评估CSV','outputs/results/matd3/baseline_exp008/curriculum_eval.csv'),
    ('最终训练摘要','outputs/results/matd3/baseline_exp008/baseline_summary.json'),
    ('冻结与复评清单','outputs/reviews/exp008_final_v1/manifest.json'),
    ('最终第5级复评','outputs/reviews/exp008_final_v1/final008_level5.json'),
    ('旧级保留对照过程','outputs/reviews/exp008_final_v1/retention_process.json'),
    ('最终复评模型视频归档','docs/baseline_exp008_final_review_delivery_v1.zip'),
    ('完整训练结束记录归档','docs/baseline_exp008_completed_records_v1.zip'),
]: link(label,path)
p('报告图由真实CSV和复评JSON重建，数据写在baseline_exp008_final_plot_data_2026-10-13.json；没有以旧曲线尾部或训练状态快照冒充最终进度。所有级别在正文按1至5展示，原文件内部level按0至4。模型、视频和评估归档以及本报告分别保留哈希，便于后续核对来源。')

h('附录 有效与保留配置的阅读提示')
p('以下几类参数容易误读。顶层target_max_speed为5、capture_radius为15并不能描述整个课程；reset实际按curriculum_schedule覆盖。本轮起始第1级对应目标1与3米每秒、保持半径50米。building_mode顶层uniform_20同样会被课程覆盖。')
p('w_terminal为240、旧w_cov为4等顶层旧模式参数不意味着baseline实际发240或重新使用覆盖。baseline分支使用baseline_success_reward为300再乘0.1；课程表旧权重也置为0。future_prediction_steps为5、escape_num_dirs为200、GAT维度等参数仍保存，但其对应预测、覆盖或图模块已关闭，不应解释为实际执行的预测长度、主动逃逸搜索或注意力层。')
p('spawn_radius和spawn_vertical_radius仍会随级更新，在008固定位置出生模式中并不控制四机到目标的环形距离；真实初距由目标出生盒、安全修正和450米最小距离检查共同决定。sigma和通用learning_rate字段不覆盖实际探索noise_scale及分开的Actor和Critic学习率。')
p('reward分项日志保留r_distance和r_teammate零列有利于比较，不能看到字段名就认为该项已生效；同理空覆盖日志或无有效路线预测值不构成策略已经学到覆盖。完整解析config.json、冻结源文件和实际统计是判断的依据。')

# Output generation uses the same blocks as the Markdown, keeping numerical
# results and narrative synchronized between editable Word and source text.
md = []
doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(.72)
section.bottom_margin = Inches(.72)
section.left_margin = Inches(.85)
section.right_margin = Inches(.85)
section.header_distance = Inches(.28)
section.footer_distance = Inches(.3)

def fontstyle(style, size, bold=False):
    style.font.name = 'Calibri'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0,0,0)
    style.element.get_or_add_rPr().append(OxmlElement('w:rFonts'))
    style.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    style.element.rPr.rFonts.set(qn('w:ascii'),'Calibri')
    style.element.rPr.rFonts.set(qn('w:hAnsi'),'Calibri')

fontstyle(doc.styles['Normal'],11)
doc.styles['Normal'].paragraph_format.space_after = Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing = 1.14
fontstyle(doc.styles['Title'],22,True)
doc.styles['Title'].paragraph_format.space_after = Pt(12)
for name,size in [('Heading 1',16),('Heading 2',12.5)]:
    fontstyle(doc.styles[name],size,True)
    doc.styles[name].paragraph_format.space_before = Pt(12)
    doc.styles[name].paragraph_format.space_after = Pt(6)
    doc.styles[name].paragraph_format.keep_with_next = True
fontstyle(doc.styles['Caption'],9)
doc.styles['Caption'].paragraph_format.space_after = Pt(8)
for name in ['Header','Footer']:
    fontstyle(doc.styles[name],8)
header = section.header.paragraphs[0]
header.text = '实验008  设置与结果总结'
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('第 ')
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)
footer.add_run(' 页')

transform = etree.XSLT(etree.parse('C:/Program Files/Microsoft Office/root/Office16/MML2OMML.XSL'))

def add_hyperlink(paragraph,label,url):
    rel = paragraph.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    element = OxmlElement('w:hyperlink'); element.set(qn('r:id'),rel)
    run = OxmlElement('w:r'); props = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'),'1764AB'); props.append(color)
    under = OxmlElement('w:u'); under.set(qn('w:val'),'single'); props.append(under)
    run.append(props); txt = OxmlElement('w:t'); txt.text = label; run.append(txt)
    element.append(run); paragraph._p.append(element)

def set_cell(cell,text,header=False,shade=False,number=False):
    cell.text = str(text)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    props = cell._tc.get_or_add_tcPr()
    margins = OxmlElement('w:tcMar')
    for side in ['top','left','bottom','right']:
        part = OxmlElement('w:'+side); part.set(qn('w:w'),'85'); part.set(qn('w:type'),'dxa'); margins.append(part)
    props.append(margins)
    if header or shade:
        fill = OxmlElement('w:shd'); fill.set(qn('w:fill'),'25465F' if header else 'F1F5F8'); props.append(fill)
    for para in cell.paragraphs:
        para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.line_spacing = 1.05
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER if number else WD_ALIGN_PARAGRAPH.LEFT
        for run in para.runs:
            run.font.size = Pt(9)
            run.font.bold = header
            if header: run.font.color.rgb = RGBColor(255,255,255)

for block in blocks:
    kind = block[0]
    if kind == 'title':
        doc.add_paragraph(block[1],'Title'); md.append('# '+block[1])
    elif kind == 'h':
        doc.add_heading(block[2],level=block[1]); md.append('#'*(block[1]+1)+' '+block[2])
    elif kind == 'p':
        doc.add_paragraph(block[1]); md.append(block[1])
    elif kind == 'eq':
        para = doc.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_after = Pt(8)
        math = etree.fromstring(('<math xmlns="http://www.w3.org/1998/Math/MathML">'+block[2]+'</math>').encode())
        omml = transform(math).getroot()
        para._p.append(omml)
        md.append('$$\n'+block[1]+'\n$$')
    elif kind == 'table':
        headers,records,widths = block[1:]
        tab = doc.add_table(rows=1, cols=len(headers))
        tab.alignment = WD_TABLE_ALIGNMENT.CENTER
        tab.autofit = False
        if widths:
            for col,w in zip(tab.columns,widths): col.width = Inches(w)
        props = tab._tbl.tblPr
        borders = OxmlElement('w:tblBorders')
        for name in ['top','left','bottom','right','insideH','insideV']:
            border=OxmlElement('w:'+name);border.set(qn('w:val'),'single');border.set(qn('w:sz'),'4');border.set(qn('w:color'),'D9D9D9');borders.append(border)
        props.append(borders)
        repeat = OxmlElement('w:tblHeader'); tab.rows[0]._tr.get_or_add_trPr().append(repeat)
        for j,text in enumerate(headers):
            set_cell(tab.rows[0].cells[j],text,header=True,number=True)
            if widths: tab.rows[0].cells[j].width = Inches(widths[j])
        for i,record in enumerate(records):
            rr=tab.add_row(); cant=OxmlElement('w:cantSplit');rr._tr.get_or_add_trPr().append(cant)
            for j,value in enumerate(record):
                value=str(value)
                numeric = len(value)<17 and (any(ch.isdigit() for ch in value) or value in ['true','false'])
                set_cell(rr.cells[j],value,shade=(i%2==1),number=numeric)
                if widths: rr.cells[j].width = Inches(widths[j])
        gap = doc.add_paragraph();gap.paragraph_format.space_after=Pt(2);gap.paragraph_format.space_before=Pt(0);gap.paragraph_format.line_spacing=Pt(2)
        md.append('| '+' | '.join(headers)+' |\n| '+' | '.join(['---']*len(headers))+' |\n'+'\n'.join('| '+' | '.join(str(v).replace('|','\\|') for v in r)+' |' for r in records))
    elif kind == 'image':
        para=doc.add_paragraph();para.paragraph_format.keep_with_next=True
        para.add_run().add_picture(str(DOCS/block[1]),width=Inches(block[3]))
        doc.add_paragraph(block[2],'Caption')
        path=str((DOCS/block[1]).resolve()).replace('\\','/')
        md.append('!['+block[2]+']('+path+')')
    elif kind == 'code':
        for line in block[1].splitlines():
            para=doc.add_paragraph();para.paragraph_format.space_after=Pt(2)
            run=para.add_run(line);run.font.name='Consolas';run.font.size=Pt(8)
        md.append('```bash\n'+block[1]+'\n```')
    elif kind in ['link','web']:
        para=doc.add_paragraph();para.paragraph_format.space_after=Pt(4)
        label,url=block[1:]
        add_hyperlink(para,label,('file:///'+url) if kind=='link' else url)
        md.append('['+label+']('+url+')')

doc.core_properties.title='无人机追捕实验008设置与结果总结'
doc.core_properties.subject='实验设置 数学实现 最终复评 旧级保留与待做事项'
doc.core_properties.author='UAV项目实验记录'
doc.core_properties.keywords='MATD3 课程学习 奖励 运动学 实验008'
for container in [doc.styles.element, doc.element]:
    for border in list(container.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.save(DOCS/(STEM+'.docx'))
(DOCS/(STEM+'.md')).write_text('\n\n'.join(md)+'\n',encoding='utf-8')
(DOCS/(STEM+'.json')).write_text(json.dumps({
    'experiment':'008','environment_steps':3000000,'critic_updates':187244,'actor_updates':93622,
    'completed_episodes':len(episodes),'completed_episode_transitions':complete,'ongoing_episode_transitions':3000000-complete,
    'stage_training_counts':stage_stats,'final_seed':60001,
    'final_metrics':[r['metrics'] for r in finals],
    'initial_new_layout_metrics':[r['metrics'] for r in initials],
    'retention_metrics':[r['metrics'] for r in retained],
    'runtime_config':config,'source_metadata':meta,'frozen_review_manifest':manifest,
    'last100_training_means':means,'decision_dt_applied':1.0,'decision_dt_accepted_next':.5,
    'monitoring_status':'PAUSED','next_training_started':False,
    'videos':[{'directory':v[0],'manifest':v[1]} for v in videos],
    'deliverables':{ext:hashlib.sha256((DOCS/(STEM+ext)).read_bytes()).hexdigest() for ext in ['.docx','.md']},
},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'docx':str(DOCS/(STEM+'.docx')),'md':str(DOCS/(STEM+'.md')),'blocks':len(blocks)},ensure_ascii=False))
